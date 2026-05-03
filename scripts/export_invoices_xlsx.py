#!/usr/bin/env python3
"""Jednorazowy eksport faktur do XLSX.

Tworzy osobne arkusze:
    - invoices
    - invoice_items
    - validation_errors (opcjonalnie, gdy wykryto braki danych)

Użycie (z poziomu głównego katalogu projektu):
        .venv/bin/python scripts/export_invoices_xlsx.py
        .venv/bin/python scripts/export_invoices_xlsx.py --out /tmp/faktury.xlsx
        .venv/bin/python scripts/export_invoices_xlsx.py --direction sale
        .venv/bin/python scripts/export_invoices_xlsx.py --month 2026-04
        .venv/bin/python scripts/export_invoices_xlsx.py --strict
"""
from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, func, select

# Upewnij się, że katalog główny projektu jest na ścieżce
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app.persistence.db import session_scope  # noqa: E402
from app.core.amount_formatting import EXCEL_NUMBER_FORMAT, EXCEL_PLN_NUMBER_FORMAT  # noqa: E402
from app.persistence.models.bank_transaction import BankTransactionORM  # noqa: E402
from app.persistence.models.invoice import InvoiceORM  # noqa: E402
from app.persistence.models.payment_allocation import PaymentAllocationORM  # noqa: E402

TWOPLACES = Decimal("0.01")
MISSING = "[MISSING]"


@dataclass(slots=True)
class ValidationErrorRow:
    invoice_id: str
    field: str
    problem: str


@dataclass(slots=True)
class PaymentStats:
    paid_amount: Decimal = Decimal("0.00")
    payment_date: date | None = None

# ---------------------------------------------------------------------------
INVOICE_COLUMNS = [
    "invoice_id",
    "direction",
    "status",
    "payment_status",
    "issue_date",
    "number_local",
    "seller_name",
    "seller_nip",
    "seller_address",
    "buyer_name",
    "buyer_nip",
    "buyer_address",
    "sale_date",
    "due_date",
    "payment_date",
    "currency",
    "total_net",
    "total_vat",
    "total_gross",
    "paid_amount",
    "remaining_amount",
    "ksef_reference_number",
    "created_at",
]

ITEM_COLUMNS = [
    "item_id",
    "invoice_id",
    "sort_order",
    "name",
    "quantity",
    "unit",
    "unit_price_net",
    "discount_amount",
    "net_amount",
    "vat_rate_pct",
    "vat_amount",
    "gross_amount",
]

REQUIRED_INVOICE_FIELDS = [
    "issue_date",
    "number_local",
    "seller_name",
    "buyer_name",
    "seller_nip",
    "buyer_nip",
    "sale_date",
    "total_net",
    "total_vat",
    "total_gross",
]

REQUIRED_ITEM_FIELDS = [
    "name",
    "quantity",
    "unit",
    "unit_price_net",
    "net_amount",
    "vat_rate_pct",
    "vat_amount",
    "gross_amount",
]

INVOICE_MONEY_COLUMNS = {
    "total_net",
    "total_vat",
    "total_gross",
    "paid_amount",
    "remaining_amount",
}

ITEM_MONEY_COLUMNS = {
    "unit_price_net",
    "discount_amount",
    "net_amount",
    "vat_amount",
    "gross_amount",
}

# ---------------------------------------------------------------------------

HEADER_FILL   = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT   = Font(color="FFFFFF", bold=True, size=10)
SALE_FILL     = PatternFill("solid", fgColor="E8F5E9")
PURCHASE_FILL = PatternFill("solid", fgColor="E3F2FD")
ERR_FILL      = PatternFill("solid", fgColor="FDE2E2")


def _as_decimal(value: object, default: Decimal = Decimal("0.00")) -> Decimal:
    if value is None:
        return default
    try:
        return Decimal(str(value))
    except Exception:  # noqa: BLE001
        return default


def _round_money(value: Decimal) -> Decimal:
    return value.quantize(TWOPLACES, rounding=ROUND_HALF_UP)


def _money_from_totals(invoice: InvoiceORM, key: str) -> Decimal:
    totals = invoice.totals_json or {}
    return _round_money(_as_decimal(totals.get(key), Decimal("0.00")))


def _snapshot_address(snapshot: dict | None) -> str:
    snap = snapshot or {}
    street = (snap.get("street") or "").strip()
    building_no = (snap.get("building_no") or "").strip()
    apartment_no = (snap.get("apartment_no") or "").strip()
    postal_code = (snap.get("postal_code") or "").strip()
    city = (snap.get("city") or "").strip()
    country = (snap.get("country") or "").strip()

    address_line = " ".join(part for part in [street, building_no] if part)
    if apartment_no:
        address_line = f"{address_line}/{apartment_no}" if address_line else apartment_no

    locality = " ".join(part for part in [postal_code, city] if part)
    composed = ", ".join(part for part in [address_line, locality, country] if part)
    return composed


def _mark_missing(value: object) -> str:
    text = "" if value is None else str(value).strip()
    return text if text else MISSING


def _is_missing(value: object) -> bool:
    if value is None:
        return True
    text = str(value).strip()
    return not text or text == MISSING


def _decimal_to_excel(value: Decimal) -> float:
    return float(_round_money(value))


def _money_number_format_for_currency(currency: object) -> str:
    normalized = str(currency or "PLN").strip().upper()
    return EXCEL_PLN_NUMBER_FORMAT if normalized == "PLN" else EXCEL_NUMBER_FORMAT


def _build_payment_stats(session, invoice_ids: list) -> dict[str, PaymentStats]:
    if not invoice_ids:
        return {}

    stmt = (
        select(
            PaymentAllocationORM.invoice_id,
            func.coalesce(func.sum(PaymentAllocationORM.allocated_amount), 0),
            func.max(BankTransactionORM.transaction_date),
        )
        .outerjoin(
            BankTransactionORM,
            BankTransactionORM.id == PaymentAllocationORM.transaction_id,
        )
        .where(
            PaymentAllocationORM.invoice_id.in_(invoice_ids),
            PaymentAllocationORM.is_reversed.is_(False),
        )
        .group_by(PaymentAllocationORM.invoice_id)
    )

    out: dict[str, PaymentStats] = {}
    for invoice_id, paid_amount, payment_date in session.execute(stmt).all():
        out[str(invoice_id)] = PaymentStats(
            paid_amount=_round_money(_as_decimal(paid_amount)),
            payment_date=payment_date,
        )
    return out


def _write_header(ws, columns: list[str]) -> None:
    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)


def _autofit(ws, columns: list[str]) -> None:
    for col_idx, col_name in enumerate(columns, start=1):
        letter = get_column_letter(col_idx)
        max_len = len(col_name) + 2
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max_len + 2, 50)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _validate_invoice_row(invoice_id: str, row: dict[str, object]) -> list[ValidationErrorRow]:
    errors: list[ValidationErrorRow] = []
    for field in REQUIRED_INVOICE_FIELDS:
        if _is_missing(row.get(field)):
            errors.append(ValidationErrorRow(invoice_id=invoice_id, field=field, problem="missing required value"))
    return errors


def _validate_item_row(invoice_id: str, row: dict[str, object]) -> list[ValidationErrorRow]:
    errors: list[ValidationErrorRow] = []
    for field in REQUIRED_ITEM_FIELDS:
        if _is_missing(row.get(field)):
            errors.append(ValidationErrorRow(invoice_id=invoice_id, field=f"item.{field}", problem="missing required value"))
    return errors


def _invoice_row(invoice: InvoiceORM, payment_stats: PaymentStats) -> dict[str, object]:
    seller_snapshot = invoice.seller_snapshot_json or {}
    buyer_snapshot = invoice.buyer_snapshot_json or {}

    total_net = _money_from_totals(invoice, "total_net")
    total_vat = _money_from_totals(invoice, "total_vat")
    total_gross = _money_from_totals(invoice, "total_gross")
    paid_amount = _round_money(payment_stats.paid_amount)
    remaining_amount = _round_money(total_gross - paid_amount)
    if remaining_amount < Decimal("0"):
        remaining_amount = Decimal("0.00")

    return {
        "invoice_id": str(invoice.id),
        "direction": (invoice.direction or "sale").strip().lower(),
        "status": invoice.status,
        "payment_status": invoice.payment_status,
        "issue_date": invoice.issue_date,
        "number_local": _mark_missing(invoice.number_local),
        "seller_name": _mark_missing(seller_snapshot.get("name")),
        "seller_nip": _mark_missing(seller_snapshot.get("nip")),
        "seller_address": _mark_missing(_snapshot_address(seller_snapshot)),
        "buyer_name": _mark_missing(buyer_snapshot.get("name")),
        "buyer_nip": _mark_missing(buyer_snapshot.get("nip")),
        "buyer_address": _mark_missing(_snapshot_address(buyer_snapshot)),
        "sale_date": invoice.sale_date,
        "due_date": invoice.due_date,
        "payment_date": payment_stats.payment_date,
        "currency": invoice.currency,
        "total_net": _decimal_to_excel(total_net),
        "total_vat": _decimal_to_excel(total_vat),
        "total_gross": _decimal_to_excel(total_gross),
        "paid_amount": _decimal_to_excel(paid_amount),
        "remaining_amount": _decimal_to_excel(remaining_amount),
        "ksef_reference_number": invoice.ksef_reference_number or "",
        "created_at": invoice.created_at,
    }


def _item_discount_amount(quantity: Decimal, unit_price_net: Decimal, net_amount: Decimal) -> Decimal:
    gross_net = _round_money(quantity * unit_price_net)
    discount = _round_money(gross_net - net_amount)
    if discount < Decimal("0"):
        return Decimal("0.00")
    return discount


def _item_row(invoice_id: str, item) -> dict[str, object]:
    quantity = _round_money(_as_decimal(item.quantity, Decimal("0")))
    unit_price_net = _round_money(_as_decimal(item.unit_price_net, Decimal("0")))
    net_amount = _round_money(_as_decimal(item.net_amount, Decimal("0")))
    vat_rate_pct = _round_money(_as_decimal(item.vat_rate, Decimal("0")))
    vat_amount = _round_money(_as_decimal(item.vat_amount, Decimal("0")))
    gross_amount = _round_money(_as_decimal(item.gross_amount, Decimal("0")))

    if vat_rate_pct == Decimal("0.00"):
        vat_amount = Decimal("0.00")

    return {
        "item_id": str(item.id),
        "invoice_id": invoice_id,
        "sort_order": item.sort_order,
        "name": _mark_missing(item.name),
        "quantity": _decimal_to_excel(quantity),
        "unit": _mark_missing(item.unit),
        "unit_price_net": _decimal_to_excel(unit_price_net),
        "discount_amount": _decimal_to_excel(_item_discount_amount(quantity, unit_price_net, net_amount)),
        "net_amount": _decimal_to_excel(net_amount),
        "vat_rate_pct": _decimal_to_excel(vat_rate_pct),
        "vat_amount": _decimal_to_excel(vat_amount),
        "gross_amount": _decimal_to_excel(gross_amount),
    }


def build_workbook(invoices: list[InvoiceORM], payment_map: dict[str, PaymentStats], strict: bool = False) -> tuple[Workbook, list[ValidationErrorRow], int]:
    wb = Workbook()
    validation_errors: list[ValidationErrorRow] = []

    # --- Arkusz 1: invoices ---
    ws_inv = wb.active
    ws_inv.title = "invoices"
    ws_inv.row_dimensions[1].height = 20
    _write_header(ws_inv, INVOICE_COLUMNS)

    for row_idx, inv in enumerate(invoices, start=2):
        fill = SALE_FILL if inv.direction == "sale" else PURCHASE_FILL
        inv_row = _invoice_row(inv, payment_map.get(str(inv.id), PaymentStats()))
        number_format = _money_number_format_for_currency(inv_row.get("currency"))
        validation_errors.extend(_validate_invoice_row(str(inv.id), inv_row))
        for col_idx, col_name in enumerate(INVOICE_COLUMNS, start=1):
            val = inv_row.get(col_name)
            if isinstance(val, date):
                val = val.isoformat()
            cell = ws_inv.cell(row=row_idx, column=col_idx, value=val)
            if col_name in INVOICE_MONEY_COLUMNS and isinstance(val, (int, float)):
                cell.number_format = number_format
            cell.fill = fill
    _autofit(ws_inv, INVOICE_COLUMNS)

    # --- Arkusz 2: invoice_items ---
    ws_items = wb.create_sheet("invoice_items")
    ws_items.row_dimensions[1].height = 20
    _write_header(ws_items, ITEM_COLUMNS)

    # Zbierz pozycje w kolejności faktur
    item_row = 2
    for inv in invoices:
        items = sorted(inv.items, key=lambda it: it.sort_order)
        fill = SALE_FILL if inv.direction == "sale" else PURCHASE_FILL
        number_format = _money_number_format_for_currency(inv.currency)
        for it in items:
            row_data = _item_row(str(inv.id), it)
            validation_errors.extend(_validate_item_row(str(inv.id), row_data))
            for col_idx, col_name in enumerate(ITEM_COLUMNS, start=1):
                val = row_data.get(col_name)
                if isinstance(val, date):
                    val = val.isoformat()
                cell = ws_items.cell(row=item_row, column=col_idx, value=val)
                if col_name in ITEM_MONEY_COLUMNS and isinstance(val, (int, float)):
                    cell.number_format = number_format
                cell.fill = fill
            item_row += 1
    _autofit(ws_items, ITEM_COLUMNS)

    if validation_errors:
        ws_err = wb.create_sheet("validation_errors")
        err_columns = ["invoice_id", "field", "problem"]
        _write_header(ws_err, err_columns)
        for row_idx, err in enumerate(validation_errors, start=2):
            ws_err.cell(row=row_idx, column=1, value=err.invoice_id).fill = ERR_FILL
            ws_err.cell(row=row_idx, column=2, value=err.field).fill = ERR_FILL
            ws_err.cell(row=row_idx, column=3, value=err.problem).fill = ERR_FILL
        _autofit(ws_err, err_columns)

    if strict and validation_errors:
        first = validation_errors[0]
        raise ValueError(
            "Walidacja eksportu nie powiodła się. "
            f"Pierwszy błąd: invoice_id={first.invoice_id}, field={first.field}, problem={first.problem}."
        )

    return wb, validation_errors, item_row - 2


def main() -> None:
    parser = argparse.ArgumentParser(description="Eksport faktur do XLSX")
    parser.add_argument("--out", default="invoices_export.xlsx",
                        help="Ścieżka pliku wyjściowego (domyślnie: invoices_export.xlsx)")
    parser.add_argument("--direction", choices=["sale", "purchase"],
                        help="Filtruj po kierunku: sale | purchase")
    parser.add_argument("--month", metavar="YYYY-MM",
                        help="Filtruj po miesiącu wystawienia, np. 2026-04")
    parser.add_argument("--strict", action="store_true",
                        help="Przerwij eksport, gdy walidacja wykryje braki")
    args = parser.parse_args()

    out_path = Path(args.out)

    month_start: date | None = None
    month_end:   date | None = None
    if args.month:
        try:
            y, m = args.month.split("-")
            month_start = date(int(y), int(m), 1)
            # Pierwszy dzień kolejnego miesiąca
            if int(m) == 12:
                month_end = date(int(y) + 1, 1, 1)
            else:
                month_end = date(int(y), int(m) + 1, 1)
        except ValueError:
            sys.exit("Błąd: --month musi być w formacie YYYY-MM, np. 2026-04")

    print(f"Łączenie z bazą danych...")
    with session_scope() as session:
        stmt = select(InvoiceORM).order_by(InvoiceORM.issue_date.asc(), InvoiceORM.created_at.asc())
        if args.direction:
            stmt = stmt.where(InvoiceORM.direction == args.direction)
        if month_start and month_end:
            stmt = stmt.where(InvoiceORM.issue_date >= month_start, InvoiceORM.issue_date < month_end)

        invoices = list(session.execute(stmt).scalars())
        # Eager-load items while session is open
        for inv in invoices:
            _ = inv.items  # trigger lazy load

        payment_map = _build_payment_stats(session, [inv.id for inv in invoices])

        print(f"Znaleziono {len(invoices)} faktur.")
        print(f"Znaleziono dane płatności dla {len(payment_map)} faktur.")

        wb, validation_errors, total_items = build_workbook(invoices, payment_map, strict=args.strict)

    wb.save(out_path)
    print(f"\nEksport zapisany: {out_path.resolve()}")
    print(f"\nArkusz 'invoices'      — {len(invoices)} wierszy, {len(INVOICE_COLUMNS)} kolumn:")
    print("  " + ", ".join(INVOICE_COLUMNS))
    print(f"\nArkusz 'invoice_items' — {total_items} wierszy, {len(ITEM_COLUMNS)} kolumn:")
    print("  " + ", ".join(ITEM_COLUMNS))
    if validation_errors:
        print(f"\nArkusz 'validation_errors' — {len(validation_errors)} wierszy.")
        print("  Eksport zawiera brakujące pola oznaczone jako [MISSING].")
    else:
        print("\nWalidacja: brak błędów.")


if __name__ == "__main__":
    main()
